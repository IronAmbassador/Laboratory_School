#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
读取模板文档内容
"""

from docx import Document

def analyze_template():
    """分析模板文档"""
    template_path = '/workspace/4-中小学辅导机构管理系统开发技术报告.docx'
    doc = Document(template_path)
    
    print("=== 模板文档段落分析 ===")
    for i, para in enumerate(doc.paragraphs[:100]):  # 只看前100个段落
        if para.text.strip():
            print(f"[{i}] {para.text[:100]}")
    
    print("\n=== 模板文档表格分析 ===")
    for i, table in enumerate(doc.tables):
        print(f"\n表格 {i}:")
        for row_idx, row in enumerate(table.rows):
            row_text = [cell.text.strip() for cell in row.cells]
            print(f"  行 {row_idx}: {row_text}")

if __name__ == '__main__':
    analyze_template()
