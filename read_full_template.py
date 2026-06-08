#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
读取模板文档的完整内容并输出结构
"""

from docx import Document

def read_full_template():
    """读取模板文档的完整内容"""
    template_path = '/workspace/4-中小学辅导机构管理系统开发技术报告.docx'
    doc = Document(template_path)
    
    print("=== 模板文档完整内容 ===")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            indent = len(para._element.xpath('./ancestor::w:pPr/w:ind'))
            print(f"[{i}] {'  '*indent}{para.text}")
    
    print("\n=== 模板文档表格 ===")
    for i, table in enumerate(doc.tables):
        print(f"\n表格 {i}:")
        for row_idx, row in enumerate(table.rows):
            row_text = [cell.text.strip() for cell in row.cells]
            print(f"  行 {row_idx}: {row_text}")

if __name__ == '__main__':
    read_full_template()
