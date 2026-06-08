#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
实验室管理系统 - 技术开发文档生成脚本
"""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import json

def read_template():
    """读取模板文档"""
    template_path = '/workspace/4-中小学辅导机构管理系统开发技术报告.docx'
    if os.path.exists(template_path):
        print(f"正在读取模板文件: {template_path}")
        doc = Document(template_path)
        print(f"模板文档包含 {len(doc.paragraphs)} 个段落")
        print(f"模板文档包含 {len(doc.tables)} 个表格")
        return doc
    else:
        print("模板文件不存在")
        return None

def create_new_document():
    """创建新的技术开发文档"""
    doc = Document()
    
    # 标题
    title = doc.add_heading('实验室管理系统开发技术报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. 项目概述
    doc.add_heading('1. 项目概述', level=1)
    doc.add_heading('1.1 项目背景', level=2)
    doc.add_paragraph(
        '随着教育现代化的发展，高校实验室数量和规模不断扩大，实验室管理工作变得日益复杂。'
        '传统的人工管理方式存在效率低、易出错、信息不透明等问题。为了提高实验室管理的效率和规范化程度，'
        '开发一套基于Web的实验室管理系统显得尤为重要。'
    )
    
    doc.add_heading('1.2 项目目标', level=2)
    doc.add_paragraph(
        '本系统旨在为高校实验室提供一个全面、高效的信息化管理平台，实现实验室资源管理、'
        '预约管理、设备管理、课程管理等核心功能的数字化和自动化，提升实验室管理水平和服务质量。'
    )
    
    doc.add_heading('1.3 项目功能简介', level=2)
    func_list = [
        '用户管理：用户注册、登录、权限管理',
        '实验室管理：实验室信息的增删改查',
        '预约管理：实验室预约申请、审核',
        '设备管理：实验室设备信息管理',
        '维修管理：设备维修记录管理',
        '课程管理：实验室课程安排',
        '项目管理：科研项目管理',
        '人员管理：实验室人员管理',
        '开放管理：实验室开放时间管理',
        '日志管理：登录日志和操作日志'
    ]
    for func in func_list:
        doc.add_paragraph(func, style='List Number')
    
    # 2. 系统架构
    doc.add_heading('2. 系统架构', level=1)
    doc.add_heading('2.1 技术架构', level=2)
    doc.add_paragraph(
        '本系统采用前后端分离的架构设计，前端使用Vue.js框架，后端使用Spring Boot框架，数据库采用MySQL。'
    )
    
    # 添加技术架构表格
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '层次'
    hdr_cells[1].text = '技术选型'
    hdr_cells[2].text = '说明'
    
    row1 = table.rows[1].cells
    row1[0].text = '前端层'
    row1[1].text = 'Vue 2.6 + Element UI'
    row1[2].text = '用户界面展示与交互'
    
    row2 = table.rows[2].cells
    row2[0].text = '服务层'
    row2[1].text = 'Spring Boot 2.5.6'
    row2[2].text = '业务逻辑处理'
    
    row3 = table.rows[3].cells
    row3[0].text = '数据层'
    row3[1].text = 'MySQL 8.0 + MyBatis Plus'
    row3[2].text = '数据持久化'
    
    doc.add_heading('2.2 技术选型', level=2)
    
    # 前端技术栈
    doc.add_heading('2.2.1 前端技术栈', level=3)
    frontend_tech = [
        'Vue 2.6.10：渐进式JavaScript框架',
        'Element UI 2.15.14：基于Vue的组件库',
        'Vue Router 3.5.4：路由管理',
        'Vuex 3.1.0：状态管理',
        'Axios 0.18.1：HTTP请求库',
        'ECharts 5.5.1：数据可视化'
    ]
    for tech in frontend_tech:
        doc.add_paragraph(tech, style='List Number')
    
    # 后端技术栈
    doc.add_heading('2.2.2 后端技术栈', level=3)
    backend_tech = [
        'Spring Boot 2.5.6：Java应用开发框架',
        'MyBatis Plus 3.5.2：ORM框架',
        'MySQL 8.0：关系型数据库',
        'Druid 1.2.8：数据库连接池',
        'Sa-Token 1.34.0：权限认证框架',
        'JWT 4.2.1：JSON Web Token',
        'Knife4j 2.0.9：API文档',
        'EasyExcel 4.0.3：Excel导入导出',
        'MinIO 7.1.0：对象存储',
        'Hutool 5.7.22：Java工具库'
    ]
    for tech in backend_tech:
        doc.add_paragraph(tech, style='List Number')
    
    # 开发工具
    doc.add_heading('2.2.3 开发工具', level=3)
    dev_tools = [
        '开发IDE：IntelliJ IDEA / VS Code',
        '版本控制：Git',
        '容器化：Docker + Docker Compose',
        '项目构建：Maven / npm'
    ]
    for tool in dev_tools:
        doc.add_paragraph(tool, style='List Number')
    
    doc.add_heading('2.3 系统架构图', level=2)
    doc.add_paragraph('【此处插入系统架构图】')
    
    # 3. 系统功能设计
    doc.add_heading('3. 系统功能设计', level=1)
    doc.add_heading('3.1 功能模块划分', level=2)
    
    # 功能模块表格
    table = doc.add_table(rows=11, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '模块名称'
    hdr_cells[1].text = '主要功能'
    hdr_cells[2].text = '说明'
    
    modules = [
        ['用户管理', '用户注册、登录、权限管理、个人中心', '系统基础模块'],
        ['实验室管理', '实验室信息的增删改查', '管理实验室基本信息'],
        ['预约管理', '预约申请、审核、取消', '管理实验室预约'],
        ['设备管理', '设备信息管理、设备状态', '管理实验设备'],
        ['维修管理', '维修申请、维修记录', '管理设备维修'],
        ['课程管理', '课程安排、课程导入', '管理实验课程'],
        ['项目管理', '项目信息管理', '管理科研项目'],
        ['人员管理', '人员信息管理', '管理实验室人员'],
        ['开放管理', '开放时间设置', '管理实验室开放'],
        ['日志管理', '登录日志、操作日志', '系统审计']
    ]
    
    for i, module in enumerate(modules):
        row = table.rows[i+1].cells
        row[0].text = module[0]
        row[1].text = module[1]
        row[2].text = module[2]
    
    doc.add_heading('3.2 核心功能详述', level=2)
    doc.add_heading('3.2.1 实验室预约管理', level=3)
    doc.add_paragraph(
        '用户可以查看可用的实验室和时间，提交预约申请。管理员可以查看所有预约申请，'
        '进行审核操作（通过/拒绝）。预约状态包括：待审核、已通过、已拒绝、已取消。'
    )
    
    doc.add_heading('3.2.2 设备管理', level=3)
    doc.add_paragraph(
        '管理实验室的设备信息，包括设备名称、购买时间、价格、保修信息、生产厂商、设备状态等。'
        '支持设备的增删改查操作。'
    )
    
    doc.add_heading('3.2.3 课程管理', level=3)
    doc.add_paragraph(
        '管理实验室的课程安排，包括课程名称、任课教师、上课时间、下课时间、使用实验室等信息。'
        '支持Excel批量导入课程信息。'
    )
    
    doc.add_heading('3.2.4 日志管理', level=3)
    doc.add_paragraph(
        '记录用户的登录日志和操作日志，包括登录时间、IP地址、操作内容、操作时间等信息，'
        '便于系统审计和追溯。'
    )
    
    # 4. 数据库设计
    doc.add_heading('4. 数据库设计', level=1)
    doc.add_heading('4.1 数据库表结构', level=2)
    
    # 主要数据表
    tables = [
        ('lab_user', '用户表'),
        ('lab_info', '实验室表'),
        ('lab_booking', '预约表'),
        ('lab_device', '设备表'),
        ('lab_device_repair', '维修表'),
        ('lab_course', '课程表'),
        ('lab_project', '项目表'),
        ('lab_person', '人员表'),
        ('lab_open', '开放表'),
        ('lab_login_log', '登录日志表'),
        ('lab_operation_log', '操作日志表')
    ]
    
    for table_name, comment in tables:
        doc.add_heading(f'4.1.1 {comment}', level=3)
        doc.add_paragraph(f'表名：{table_name}')
    
    doc.add_heading('4.2 ER图', level=2)
    doc.add_paragraph('【此处插入ER图】')
    
    # 5. 系统实现
    doc.add_heading('5. 系统实现', level=1)
    doc.add_heading('5.1 项目结构', level=2)
    
    # 后端项目结构
    doc.add_heading('5.1.1 后端项目结构', level=3)
    backend_structure = [
        'com.laboratory/',
        '  ├── annotation/          # 自定义注解',
        '  ├── aspect/              # 切面',
        '  ├── config/              # 配置类',
        '  ├── controller/          # 控制器',
        '  ├── exception/           # 异常处理',
        '  ├── handler/             # 处理器',
        '  ├── interceptor/         # 拦截器',
        '  ├── listener/            # 监听器',
        '  ├── mapper/              # 数据访问层',
        '  ├── model/               # 模型层',
        '  │   ├── constant/        # 常量',
        '  │   ├── converter/       # 转换器',
        '  │   ├── dto/             # 数据传输对象',
        '  │   ├── entity/          # 实体类',
        '  │   ├── enums/           # 枚举',
        '  │   └── vo/              # 视图对象',
        '  ├── result/              # 统一响应',
        '  ├── service/             # 业务逻辑层',
        '  └── utils/               # 工具类'
    ]
    for line in backend_structure:
        doc.add_paragraph(line)
    
    # 前端项目结构
    doc.add_heading('5.1.2 前端项目结构', level=3)
    frontend_structure = [
        'src/',
        '  ├── api/                 # API接口',
        '  ├── assets/              # 静态资源',
        '  ├── components/          # 组件',
        '  ├── icons/               # 图标',
        '  ├── layout/              # 布局',
        '  ├── router/              # 路由',
        '  ├── store/               # 状态管理',
        '  ├── styles/              # 样式',
        '  ├── utils/               # 工具函数',
        '  └── views/               # 页面'
    ]
    for line in frontend_structure:
        doc.add_paragraph(line)
    
    doc.add_heading('5.2 核心代码实现', level=2)
    doc.add_heading('5.2.1 后端核心实现', level=3)
    
    # 启动类
    doc.add_heading('启动类', level=4)
    doc.add_paragraph('LaboratoryApplication.java - Spring Boot 启动类')
    
    # 控制器示例
    doc.add_heading('控制器示例', level=4)
    doc.add_paragraph('BookingController.java - 预约控制器')
    
    # 服务层示例
    doc.add_heading('服务层示例', level=4)
    doc.add_paragraph('BookingService.java - 预约服务接口')
    doc.add_paragraph('BookingServiceImpl.java - 预约服务实现')
    
    doc.add_heading('5.2.2 前端核心实现', level=3)
    
    # 主入口
    doc.add_heading('主入口', level=4)
    doc.add_paragraph('main.js - Vue 应用入口')
    
    # 路由配置
    doc.add_heading('路由配置', level=4)
    doc.add_paragraph('router/index.js - 路由配置')
    
    # 6. 系统部署
    doc.add_heading('6. 系统部署', level=1)
    doc.add_heading('6.1 部署环境', level=2)
    deploy_env = [
        '操作系统：Linux / Windows',
        'JDK版本：17',
        'Node版本：>= 8.9',
        '数据库：MySQL 8.0',
        '容器：Docker + Docker Compose'
    ]
    for env in deploy_env:
        doc.add_paragraph(env, style='List Number')
    
    doc.add_heading('6.2 部署步骤', level=2)
    deploy_steps = [
        '1. 克隆项目代码到本地',
        '2. 配置数据库连接信息',
        '3. 执行数据库初始化脚本',
        '4. 使用 Docker Compose 启动服务',
        '5. 访问系统进行测试'
    ]
    for step in deploy_steps:
        doc.add_paragraph(step)
    
    doc.add_heading('6.3 Docker Compose 配置', level=2)
    doc.add_paragraph('使用 docker-compose.yml 配置文件，一键启动 MySQL、后端服务、前端服务。')
    
    # 7. 系统测试
    doc.add_heading('7. 系统测试', level=1)
    doc.add_heading('7.1 测试环境', level=2)
    doc.add_paragraph('测试环境与部署环境保持一致。')
    
    doc.add_heading('7.2 功能测试', level=2)
    test_items = [
        '用户注册、登录功能测试',
        '实验室管理功能测试',
        '预约申请与审核功能测试',
        '设备管理功能测试',
        '课程管理功能测试',
        '日志查询功能测试'
    ]
    for item in test_items:
        doc.add_paragraph(item, style='List Number')
    
    doc.add_heading('7.3 测试结果', level=2)
    doc.add_paragraph('经过全面测试，系统各项功能运行正常，满足设计要求。')
    
    # 8. 总结与展望
    doc.add_heading('8. 总结与展望', level=1)
    doc.add_heading('8.1 项目总结', level=2)
    doc.add_paragraph(
        '本项目成功实现了一个功能完善的实验室管理系统，采用前后端分离架构，'
        '使用了当前主流的技术栈，具有良好的用户体验和可维护性。'
    )
    
    doc.add_heading('8.2 不足之处', level=2)
    doc.add_paragraph(
        '系统在大数据量下的性能优化、移动端适配等方面还有提升空间。'
    )
    
    doc.add_heading('8.3 未来展望', level=2)
    future = [
        '增加数据分析和可视化功能',
        '开发移动端APP',
        '集成AI智能预约推荐',
        '优化系统性能',
        '增加更多的安全措施'
    ]
    for item in future:
        doc.add_paragraph(item, style='List Number')
    
    # 保存文档
    output_path = '/workspace/实验室管理系统开发技术报告.docx'
    doc.save(output_path)
    print(f"文档已生成: {output_path}")
    return output_path

if __name__ == '__main__':
    # 先尝试读取模板
    template_doc = read_template()
    
    # 生成新文档
    output_file = create_new_document()
    print(f"技术开发文档生成完成！")
